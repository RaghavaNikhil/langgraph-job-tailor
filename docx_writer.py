"""Writes the cover letter to a clean .docx business-letter document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt


def write_cover_letter_docx(
    letter_text: str,
    output_path: Path,
    applicant_name: str,
    contact_line: str = "",
) -> Path:
    """Write the letter body into a .docx with a simple contact header."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    header = doc.add_paragraph()
    run = header.add_run(applicant_name)
    run.bold = True
    run.font.size = Pt(14)
    if contact_line:
        doc.add_paragraph(contact_line)
    doc.add_paragraph(date.today().strftime("%B %d, %Y"))
    doc.add_paragraph("")

    for block in letter_text.split("\n\n"):
        block = block.strip()
        if block:
            doc.add_paragraph(block)

    try:
        doc.save(str(output_path))
        return output_path
    except PermissionError:
        for n in range(2, 20):
            alt = output_path.with_stem(f"{output_path.stem}_{n}")
            try:
                doc.save(str(alt))
                print(
                    f"[docx] {output_path.name} is locked. "
                    f"Saved as {alt.name} instead."
                )
                return alt
            except PermissionError:
                continue
        raise
